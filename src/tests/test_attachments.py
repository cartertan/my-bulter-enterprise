"""Test mission attachment isolation, image boundaries and Office edit fidelity.

Inputs: Synthetic files and isolated clients. Outputs: assertions, no external
API calls or private data. Author: Carter Tan.
"""
import base64
import io
import json
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from fastapi.testclient import TestClient

from src.app import create_app
from src.attachments import parse, template_output, edited_output, office_paragraphs
from src.providers import request_spec


def office(kind: str, text: str) -> bytes:
    """Build a real Office fixture with styled source text."""
    output = io.BytesIO()
    if kind == 'docx':
        document = Document()
        document.add_paragraph(text)
    else:
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        from pptx.util import Inches
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text = text
    document.save(output)
    return output.getvalue()


def completed(output: str) -> dict:
    """Provide the minimal completed mission shape."""
    return {'status': 'completed', 'prompt': 'Synthetic request', 'steps': [{'output': output}]}


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_edit_and_template_preserve_parts(kind: str) -> None:
    """Unchanged package parts retain their exact bytes; replacement is grounded."""
    source = office(kind, 'Original example')
    result = edited_output(source, 'sample.' + kind, completed(json.dumps({'replacements': [{'old': 'Original example', 'new': 'New approved wording'}]})))
    assert office_paragraphs(result, 'result.' + kind) == ['New approved wording']
    with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(io.BytesIO(result)) as new:
        for name in old.namelist():
            if name != 'word/document.xml' and not name.startswith('ppt/slides/slide'):
                assert old.read(name) == new.read(name)
    template = office(kind, '{{RESPONSE}}')
    assert 'Completed text' in office_paragraphs(template_output(template, 'sample.' + kind, completed('Completed text')), 'result.' + kind)
    with pytest.raises(ValueError, match='invalid or ambiguous'):
        edited_output(source, 'sample.' + kind, completed('{"replacements":[{"old":"absent","new":"bad"}]}'))


def test_photo_normalized_and_limits() -> None:
    """Images are decoded, resized and re-encoded, with malformed inputs refused."""
    output = io.BytesIO()
    Image.new('RGB', (2000, 1000), 'red').save(output, 'PNG')
    parsed = parse(output.getvalue(), 'photo.png')
    assert parsed['width'] == 1600 and parsed['height'] == 800
    assert base64.b64decode(parsed['image']).startswith(b'\xff\xd8')
    assert request_spec('ollama', 'vision', 'system', [{'role': 'user', 'content': 'Read photo', 'images': [parsed['image']]}])[2]['messages'][1]['images']
    with pytest.raises(ValueError):
        parse(b'x' * 12001, 'long.txt')
    with pytest.raises(Exception):
        parse(b'not an image', 'fake.png')


def test_attachment_mission_flow(tmp_path: Path, monkeypatch) -> None:
    """Selected references reach inference, persist on the mission, never become memory."""
    observed = []
    async def fake_generate(provider, model, system, messages, **kwargs):
        observed.append(messages)
        yield 'Synthetic analysis'
    monkeypatch.setattr('src.app.generate', fake_generate)
    app = create_app(tmp_path / 'test.sqlite3')
    with TestClient(app) as client:
        assert client.post('/api/attachments?filename=test.txt', content=b'test').status_code == 403
        client.headers['x-bulter-token'] = app.state.csrf
        result = client.post('/api/attachments?filename=sample.txt', content=b'Evidence uniquely attached to this mission')
        assert result.status_code == 200, result.text
        attachment = result.json()
        assert 'raw' not in attachment and 'image' not in attachment
        assert client.post('/api/attachments?filename=../bad.txt', content=b'x').status_code == 400
        bot = client.post('/api/bulters', json={'name':'Test', 'role':'Analysis', 'instructions':'Use evidence', 'provider':'ollama','model':'test','shared_memory':False}).json()
        response = client.post('/api/missions', json={'prompt':'Review the evidence', 'bulter_ids':[bot['id']], 'attachment_ids':[attachment['id']]})
        assert response.status_code == 200, response.text
        for _ in range(100):
            mission = app.state.store.get('mission', response.json()['id'])
            if mission['status'] not in {'queued', 'running'}:
                break
            time.sleep(.01)
        assert mission['status'] == 'completed'
        assert 'Evidence uniquely attached' in observed[0][0]['content']
        assert not app.state.store.all('memory')
        assert client.post('/api/missions', json={'prompt':'Review', 'bulter_ids':[bot['id']], 'attachment_ids':['missing']}).status_code == 404
        assert client.post('/api/missions', json={'prompt':'Review', 'bulter_ids':[bot['id']], 'attachment_mode':'edit'}).status_code == 400
        cloud_bot = client.post('/api/bulters', json={'name':'Cloud', 'role':'Analysis', 'instructions':'Use evidence', 'provider':'openai','model':'test','shared_memory':False}).json()
        response = client.post('/api/missions', json={'prompt':'Review', 'bulter_ids':[cloud_bot['id']], 'attachment_ids':[attachment['id']], 'cloud_consent':True})
        assert response.status_code == 400 and 'attachments' in response.json()['detail']


def test_photo_delivery_and_template_export(tmp_path: Path, monkeypatch) -> None:
    """Verify photo bytes reach inference and validated edits reach the download."""
    observed = []
    async def fake_vision(bots):
        return None
    async def fake_generate(provider, model, system, messages, **kwargs):
        observed.append(messages)
        yield json.dumps({'replacements':[{'old':'Original example','new':'Reviewed replacement'}]})
    monkeypatch.setattr('src.app.require_vision', fake_vision)
    monkeypatch.setattr('src.app.generate', fake_generate)
    app = create_app(tmp_path / 'test.sqlite3')
    with TestClient(app) as client:
        client.headers['x-bulter-token'] = app.state.csrf
        output = io.BytesIO()
        Image.new('RGB', (64,64), 'blue').save(output, 'PNG')
        photo = client.post('/api/attachments?filename=photo.png', content=output.getvalue()).json()
        template = client.post('/api/attachments?filename=sample.docx', content=office('docx', 'Original example')).json()
        assert client.get(f"/api/attachments/{photo['id']}/preview").content.startswith(b'\xff\xd8')
        bot = client.post('/api/bulters', json={'name':'Vision', 'role':'Analysis','instructions':'Use evidence','provider':'ollama','model':'vision','shared_memory':False}).json()
        result = client.post('/api/missions', json={'prompt':'Replace example wording', 'bulter_ids':[bot['id']], 'attachment_ids':[photo['id'],template['id']], 'attachment_mode':'edit', 'template_id':template['id']})
        assert result.status_code == 200, result.text
        for _ in range(100):
            mission = app.state.store.get('mission', result.json()['id'])
            if mission['status'] not in {'running','queued'}:
                break
            time.sleep(.01)
        assert observed[0][0]['images']
        assert 'editable_paragraphs' in observed[0][0]['content']
        response = client.get(f"/api/missions/{mission['id']}/export/docx?use_template=true")
        assert response.status_code == 200, response.text
        assert office_paragraphs(response.content, 'edited.docx') == ['Reviewed replacement']


def test_vision_capability_fails_closed(monkeypatch) -> None:
    """A text-only model is rejected before an image inference call."""
    import asyncio
    import httpx
    from fastapi import HTTPException
    from src.attachment_routes import require_vision
    original = httpx.AsyncClient
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json={'capabilities':['completion']}))
    monkeypatch.setattr('src.attachment_routes.httpx.AsyncClient', lambda **kwargs: original(transport=transport, **kwargs))
    with pytest.raises(HTTPException, match='does not support vision'):
        asyncio.run(require_vision([{'provider':'ollama','model':'text-only'}]))
    with pytest.raises(HTTPException, match='Ollama'):
        asyncio.run(require_vision([{'provider':'openai','model':'cloud'}]))
