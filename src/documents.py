"""Bounded document ingestion and editable mission exports.

Inputs: Uploaded bytes or saved mission text. Outputs: text, DOCX and PPTX bytes.
Assumptions: Local trusted user; document text is untrusted reference data.
No macros, links, or embedded files are executed or fetched. Author: Carter Tan.
"""

import io
import re
import textwrap
import zipfile
import json
import sys
import unicodedata
from pathlib import Path

MAX_UPLOAD = 10 * 1024 * 1024
MAX_TEXT = 120_000
EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".pptx"}


def extract_text(data: bytes, filename: str) -> str:
    """Extract visible text with explicit size limits; reject empty scans."""
    suffix = Path(filename).suffix.lower()
    if suffix not in EXTENSIONS or not data or len(data) > MAX_UPLOAD:
        raise ValueError("Choose a nonempty TXT, MD, PDF, DOCX or PPTX file up to 10 MB.")
    try:
        if suffix in {".docx", ".pptx"}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(m.file_size for m in members) > 40 * 1024 * 1024:
                    raise ValueError("Office document is too large after decompression.")
                if any("vbaproject" in m.filename.lower() for m in members):
                    raise ValueError("Macro content is not supported.")
        if suffix in {".txt", ".md"}:
            text = data.decode("utf-8-sig")
        elif suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(data))
            text = "\n".join(block.text if hasattr(block, "text") else
                             "\n".join(" | ".join(c.text for c in row.cells) for row in block.rows)
                             for block in document.iter_inner_content())
        elif suffix == ".pptx":
            from pptx import Presentation
            deck = Presentation(io.BytesIO(data))
            text = "\n\n".join(f"Slide {i + 1}\n" + "\n".join(
                shape.text if shape.has_text_frame else
                "\n".join(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
                for shape in slide.shapes if shape.has_text_frame or shape.has_table)
                for i, slide in enumerate(deck.slides))
        else:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as document:
                if len(document.pages) > 200:
                    raise ValueError("PDFs are limited to 200 pages.")
                pages = [page.extract_text() or "" for page in document.pages]
                text = "\n\n".join(f"Page {i + 1}\n{page}" for i, page in enumerate(pages))
                if not any(page.strip() for page in pages):
                    raise ValueError("No readable text. Scanned documents need OCR first.")
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()
        if not text or len(text) > MAX_TEXT:
            raise ValueError("Extracted text must contain 1–120,000 characters. Split large documents first.")
        return text
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Cannot read this document. Check its format, encoding and password protection.") from exc


def chunks(text: str, size: int = 2600) -> list[str]:
    """Split without dropping text; prefer paragraph boundaries."""
    result = []
    while text:
        cut = len(text) if len(text) <= size else text.rfind("\n", 0, size)
        if cut < size // 2 and len(text) > size:
            cut = size
        result.append(text[:cut].strip())
        text = text[cut:].strip()
    return [part for part in result if part]


def export_mission(mission: dict, format_name: str) -> bytes:
    """Export the final completed response; preserve wording and provenance."""
    if mission.get("status") != "completed" or not mission.get("steps"):
        raise ValueError("Exports require a completed mission.")
    step = mission["steps"][-1]
    text = step.get("output", "").strip()
    if not text:
        raise ValueError("This mission has no final response to export.")
    provenance = f"{step['name']} · {step['model']} · {mission['created']}"
    title = re.sub(r"[^\w\s]", "", mission["prompt"].splitlines()[0])[:80].strip() or "Mission response"
    output = io.BytesIO()
    if format_name == "docx":
        from docx import Document
        from docx.shared import Pt, RGBColor
        doc = Document()
        doc.core_properties.author = "My Bulter Enterprise"
        doc.core_properties.title = title
        # Remove template decoration so generated report titles remain plain black.
        for border in doc.styles.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
            doc.styles[name].font.name = "Calibri"
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_heading(title, 0)
        doc.add_paragraph(provenance)
        doc.add_heading("Request", 1)
        doc.add_paragraph(mission["prompt"])
        doc.add_heading("Response", 1)
        for line in text.splitlines():
            heading = re.match(r"^(#{1,3})\s+(.+)", line)
            if heading:
                doc.add_heading(heading[2], len(heading[1]))
            else:
                doc.add_paragraph(line)
        doc.save(output)
    elif format_name == "pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        # Bounded lines preserve every word without allowing a dense response to overflow.
        lines = [part for line in text.splitlines() for part in
                 (textwrap.wrap(line, width=42 if any(unicodedata.east_asian_width(c) in "WF" for c in line) else 72,
                                replace_whitespace=False) or [""])]
        pages = [lines[i:i + 14] for i in range(0, len(lines), 14)]
        for index, page in enumerate(pages):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string("191817")
            for x, y, w, h, content, size, color in (
                (.65, .4, 12, .6, f"Mission response  {index + 1} / {len(pages)}", 28, "D7B780"),
                (.65, 1.25, 12, 5.5, "\n".join(page), 20, "F1ECE3"),
                (.65, 7, 12, .3, provenance[:110], 10, "D7B780"),
            ):
                frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
                frame.word_wrap = True
                frame.text = content
                for paragraph in frame.paragraphs:
                    paragraph.font.size = Pt(size)
                    paragraph.font.color.rgb = RGBColor.from_string(color)
            slide.notes_slide.notes_text_frame.text = "Original request: " + mission["prompt"]
        deck.save(output)
    else:
        raise ValueError("Choose docx or pptx.")
    return output.getvalue()


if __name__ == "__main__":
    # Extraction runs in a short-lived child so a malformed parser input can time out.
    try:
        import resource
        # macOS does not reliably support RLIMIT_AS; compressed inputs have a separate cap.
        if sys.platform != "darwin":
            resource.setrlimit(resource.RLIMIT_AS, (1024 ** 3, 1024 ** 3))
        resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
        result = {"text": extract_text(sys.stdin.buffer.read(MAX_UPLOAD + 1), sys.argv[1])}
    except Exception as error:
        result = {"error": str(error) if isinstance(error, ValueError) else "Document extraction failed."}
    sys.stdout.write(json.dumps(result))
